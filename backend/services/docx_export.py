"""Renders a KYC one-pager record into a downloadable .docx file."""
import io
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def render_onepager_docx(record: dict) -> io.BytesIO:
    onepager = record["onepager"]
    doc = Document()

    title = doc.add_heading(onepager.get("company_name") or record.get("company_name") or "Company Briefing", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if onepager.get("headline"):
        p = doc.add_paragraph(onepager["headline"])
        p.runs[0].italic = True

    for section in onepager.get("sections", []):
        doc.add_heading(section.get("title", ""), level=1)
        for bullet in section.get("bullets", []):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(bullet.get("text", ""))
            sid = bullet.get("source_id")
            if sid:
                sup = p.add_run(f" [{sid}]")
                sup.font.superscript = True
                sup.font.size = Pt(9)

    if onepager.get("sources"):
        doc.add_heading("Sources", level=1)
        for s in onepager["sources"]:
            line = f"[{s['id']}] {s['label']}"
            if s.get("url"):
                line += f" — {s['url']}"
            doc.add_paragraph(line, style="List Number")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
